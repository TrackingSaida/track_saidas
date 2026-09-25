"""Importação em massa de Bases via planilha (PRD-002).

Fluxo: modelo → preview (sem gravar) → confirmar (create/update de preços).
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Sequence, Tuple

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from models import BasePreco, User

MAX_LINHAS = 200
COLUNAS = ("base", "flex", "shopee", "avulso")
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MODELO_FILENAME = "modelo-importacao-bases.xlsx"

ACAO_CRIAR = "criar"
ACAO_ATUALIZAR = "atualizar"
ACAO_ERRO = "erro"

STATUS_CRIADA = "criada"
STATUS_ATUALIZADA = "atualizada"
STATUS_ERRO = "erro"


def normalize_base_nome(nome: Optional[str]) -> str:
    return re.sub(r"\s+", " ", (nome or "").strip()).upper()


def parse_moeda(valor: Any) -> float:
    """Aceita 8, 8.5, 8,50, R$ 8,00, 1.234,56."""
    if valor is None:
        raise ValueError("Valor obrigatório.")
    if isinstance(valor, bool):
        raise ValueError("Valor inválido.")
    if isinstance(valor, (int, float)):
        num = float(valor)
        if num < 0:
            raise ValueError("Valor não pode ser negativo.")
        return num

    s = str(valor).strip()
    if not s:
        raise ValueError("Valor obrigatório.")
    s = re.sub(r"(?i)r\$\s*", "", s).strip()
    s = s.replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        num = float(s)
    except ValueError as exc:
        raise ValueError("Valor inválido.") from exc
    if num < 0:
        raise ValueError("Valor não pode ser negativo.")
    return num


def gerar_modelo_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bases"
    ws.append(list(COLUNAS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    # Linha de exemplo — usuário deve remover ou sobrescrever
    ws.append(["EXEMPLO BASE", 8.0, 5.0, 8.0])
    ws.column_dimensions["A"].width = 24
    for col in ("B", "C", "D"):
        ws.column_dimensions[col].width = 12
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _normalizar_cabecalho(valor: Any) -> str:
    return re.sub(r"\s+", " ", str(valor or "").strip()).lower()


def _mapear_colunas(header_row: Sequence[Any]) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for idx, raw in enumerate(header_row):
        key = _normalizar_cabecalho(raw)
        if key in COLUNAS and key not in mapping:
            mapping[key] = idx
    faltando = [c for c in COLUNAS if c not in mapping]
    if faltando:
        raise HTTPException(
            status_code=400,
            detail=f"Planilha sem colunas obrigatórias: {', '.join(faltando)}.",
        )
    return mapping


def ler_linhas_xlsx(content: bytes) -> List[Dict[str, Any]]:
    """Lê o XLSX e devolve linhas brutas (ainda sem classificar)."""
    if not content:
        raise HTTPException(status_code=400, detail="Arquivo vazio.")
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Arquivo inválido. Envie uma planilha .xlsx.",
        ) from exc

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise HTTPException(status_code=400, detail="Planilha sem cabeçalho.") from None

    if not header or all(c is None or str(c).strip() == "" for c in header):
        raise HTTPException(status_code=400, detail="Planilha sem cabeçalho.")

    colmap = _mapear_colunas(header)
    linhas: List[Dict[str, Any]] = []
    for offset, row in enumerate(rows_iter, start=2):
        if row is None:
            continue
        cells = list(row)
        vals = {
            "linha": offset,
            "base": cells[colmap["base"]] if colmap["base"] < len(cells) else None,
            "flex": cells[colmap["flex"]] if colmap["flex"] < len(cells) else None,
            "shopee": cells[colmap["shopee"]] if colmap["shopee"] < len(cells) else None,
            "avulso": cells[colmap["avulso"]] if colmap["avulso"] < len(cells) else None,
        }
        if all(
            v is None or (isinstance(v, str) and not v.strip())
            for k, v in vals.items()
            if k != "linha"
        ):
            continue
        linhas.append(vals)

    if len(linhas) > MAX_LINHAS:
        raise HTTPException(
            status_code=400,
            detail=f"A planilha tem mais de {MAX_LINHAS} linhas. Reduza o arquivo e tente de novo.",
        )
    return linhas


@dataclass
class LinhaImport:
    linha: int
    base: str
    flex: Optional[float]
    shopee: Optional[float]
    avulso: Optional[float]
    acao: str
    id_base: Optional[int] = None
    atual: Optional[Dict[str, float]] = None
    motivo: Optional[str] = None

    def as_preview_dict(self) -> Dict[str, Any]:
        return {
            "linha": self.linha,
            "base": self.base,
            "flex": self.flex,
            "shopee": self.shopee,
            "avulso": self.avulso,
            "acao": self.acao,
            "id_base": self.id_base,
            "atual": self.atual,
            "motivo": self.motivo,
        }


@dataclass
class PreviewResult:
    linhas: List[LinhaImport] = field(default_factory=list)

    @property
    def resumo(self) -> Dict[str, int]:
        criar = sum(1 for l in self.linhas if l.acao == ACAO_CRIAR)
        atualizar = sum(1 for l in self.linhas if l.acao == ACAO_ATUALIZAR)
        erro = sum(1 for l in self.linhas if l.acao == ACAO_ERRO)
        return {
            "criar": criar,
            "atualizar": atualizar,
            "erro": erro,
            "total_linhas": len(self.linhas),
        }

    def as_dict(self) -> Dict[str, Any]:
        return {
            "ok": True,
            "resumo": self.resumo,
            "linhas": [l.as_preview_dict() for l in self.linhas],
        }


def _existentes_por_nome(db: Session, sub_base: str) -> Dict[str, BasePreco]:
    rows = db.scalars(
        select(BasePreco).where(BasePreco.sub_base == sub_base)
    ).all()
    out: Dict[str, BasePreco] = {}
    for obj in rows:
        key = normalize_base_nome(obj.base)
        if key and key not in out:
            out[key] = obj
    return out


def classificar_linhas(
    linhas_brutas: Sequence[Dict[str, Any]],
    existentes: Dict[str, BasePreco],
) -> PreviewResult:
    """Classifica linhas em criar / atualizar / erro (sem gravar)."""
    result = PreviewResult()
    vistos: Dict[str, int] = {}

    for raw in linhas_brutas:
        n_linha = int(raw.get("linha") or 0)
        nome = normalize_base_nome(str(raw.get("base") if raw.get("base") is not None else ""))
        if not nome:
            result.linhas.append(
                LinhaImport(
                    linha=n_linha,
                    base="",
                    flex=None,
                    shopee=None,
                    avulso=None,
                    acao=ACAO_ERRO,
                    motivo="Nome obrigatório.",
                )
            )
            continue

        if nome in vistos:
            result.linhas.append(
                LinhaImport(
                    linha=n_linha,
                    base=nome,
                    flex=None,
                    shopee=None,
                    avulso=None,
                    acao=ACAO_ERRO,
                    motivo=f"Nome repetido na planilha (já na linha {vistos[nome]}).",
                )
            )
            continue

        try:
            flex = parse_moeda(raw.get("flex"))
            shopee = parse_moeda(raw.get("shopee"))
            avulso = parse_moeda(raw.get("avulso"))
        except ValueError as exc:
            result.linhas.append(
                LinhaImport(
                    linha=n_linha,
                    base=nome,
                    flex=None,
                    shopee=None,
                    avulso=None,
                    acao=ACAO_ERRO,
                    motivo=str(exc) or "Valor inválido.",
                )
            )
            continue

        vistos[nome] = n_linha
        existente = existentes.get(nome)
        if existente is not None:
            result.linhas.append(
                LinhaImport(
                    linha=n_linha,
                    base=nome,
                    flex=flex,
                    shopee=shopee,
                    avulso=avulso,
                    acao=ACAO_ATUALIZAR,
                    id_base=int(existente.id_base),
                    atual={
                        "flex": float(existente.ml or 0),
                        "shopee": float(existente.shopee or 0),
                        "avulso": float(existente.avulso or 0),
                    },
                )
            )
        else:
            result.linhas.append(
                LinhaImport(
                    linha=n_linha,
                    base=nome,
                    flex=flex,
                    shopee=shopee,
                    avulso=avulso,
                    acao=ACAO_CRIAR,
                )
            )

    return result


def preview_importacao(db: Session, sub_base: str, content: bytes) -> Dict[str, Any]:
    brutas = ler_linhas_xlsx(content)
    existentes = _existentes_por_nome(db, sub_base)
    return classificar_linhas(brutas, existentes).as_dict()


def _validar_linha_confirm(
    item: Dict[str, Any],
    existentes: Dict[str, BasePreco],
    vistos: Dict[str, int],
) -> Tuple[Optional[LinhaImport], Optional[str]]:
    n_linha = int(item.get("linha") or 0)
    nome = normalize_base_nome(str(item.get("base") if item.get("base") is not None else ""))
    if not nome:
        return None, "Nome obrigatório."
    if nome in vistos:
        return None, f"Nome repetido na planilha (já na linha {vistos[nome]})."
    try:
        flex = parse_moeda(item.get("flex"))
        shopee = parse_moeda(item.get("shopee"))
        avulso = parse_moeda(item.get("avulso"))
    except ValueError as exc:
        return None, str(exc) or "Valor inválido."

    vistos[nome] = n_linha
    existente = existentes.get(nome)
    if existente is not None:
        return (
            LinhaImport(
                linha=n_linha,
                base=nome,
                flex=flex,
                shopee=shopee,
                avulso=avulso,
                acao=ACAO_ATUALIZAR,
                id_base=int(existente.id_base),
            ),
            None,
        )
    return (
        LinhaImport(
            linha=n_linha,
            base=nome,
            flex=flex,
            shopee=shopee,
            avulso=avulso,
            acao=ACAO_CRIAR,
        ),
        None,
    )


def confirmar_importacao(
    db: Session,
    sub_base: str,
    current_user: User,
    linhas_in: Sequence[Dict[str, Any]],
) -> Dict[str, Any]:
    if not isinstance(linhas_in, (list, tuple)):
        raise HTTPException(status_code=400, detail="Envie a lista de linhas confirmadas.")
    if len(linhas_in) > MAX_LINHAS:
        raise HTTPException(
            status_code=400,
            detail=f"Máximo de {MAX_LINHAS} linhas por importação.",
        )
    if len(linhas_in) == 0:
        raise HTTPException(status_code=400, detail="Nenhuma linha para importar.")

    existentes = _existentes_por_nome(db, sub_base)
    vistos: Dict[str, int] = {}
    detalhes: List[Dict[str, Any]] = []
    criadas = 0
    atualizadas = 0
    erros = 0

    username = getattr(current_user, "username", None)

    for item in linhas_in:
        if not isinstance(item, dict):
            erros += 1
            detalhes.append(
                {
                    "linha": 0,
                    "base": "",
                    "status": STATUS_ERRO,
                    "motivo": "Linha inválida.",
                }
            )
            continue

        linha_ok, motivo = _validar_linha_confirm(item, existentes, vistos)
        n_linha = int(item.get("linha") or 0)
        if linha_ok is None:
            erros += 1
            detalhes.append(
                {
                    "linha": n_linha,
                    "base": normalize_base_nome(str(item.get("base") or "")),
                    "status": STATUS_ERRO,
                    "motivo": motivo or "Linha inválida.",
                }
            )
            continue

        try:
            if linha_ok.acao == ACAO_ATUALIZAR:
                obj = existentes.get(linha_ok.base)
                if obj is None or (obj.sub_base or "") != sub_base:
                    erros += 1
                    detalhes.append(
                        {
                            "linha": linha_ok.linha,
                            "base": linha_ok.base,
                            "status": STATUS_ERRO,
                            "motivo": "Base não encontrada nesta sub_base.",
                        }
                    )
                    continue
                obj.ml = float(linha_ok.flex)  # type: ignore[arg-type]
                obj.shopee = float(linha_ok.shopee)  # type: ignore[arg-type]
                obj.avulso = float(linha_ok.avulso)  # type: ignore[arg-type]
                # não altera ativo / agenda / nome
                db.add(obj)
                atualizadas += 1
                detalhes.append(
                    {
                        "linha": linha_ok.linha,
                        "base": linha_ok.base,
                        "status": STATUS_ATUALIZADA,
                        "id_base": int(obj.id_base),
                    }
                )
            else:
                obj = BasePreco(
                    base=linha_ok.base,
                    sub_base=sub_base,
                    username=username,
                    shopee=float(linha_ok.shopee),  # type: ignore[arg-type]
                    ml=float(linha_ok.flex),  # type: ignore[arg-type]
                    avulso=float(linha_ok.avulso),  # type: ignore[arg-type]
                    ativo=True,
                    dias_coleta=[1, 2, 3, 4, 5, 6],
                    agenda_coleta_confirmada=False,
                )
                db.add(obj)
                db.flush()
                existentes[linha_ok.base] = obj
                criadas += 1
                detalhes.append(
                    {
                        "linha": linha_ok.linha,
                        "base": linha_ok.base,
                        "status": STATUS_CRIADA,
                        "id_base": int(obj.id_base),
                    }
                )
        except Exception:
            erros += 1
            detalhes.append(
                {
                    "linha": linha_ok.linha,
                    "base": linha_ok.base,
                    "status": STATUS_ERRO,
                    "motivo": "Falha ao gravar esta linha.",
                }
            )

    db.commit()
    return {
        "ok": True,
        "criadas": criadas,
        "atualizadas": atualizadas,
        "erros": erros,
        "detalhes": detalhes,
    }
